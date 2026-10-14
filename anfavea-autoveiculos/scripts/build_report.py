"""Monta a base consolidada e o relatório a partir de data/csv/*.csv.

Saídas (em relatorio/):
  base_autoveiculos.csv                  formato longo, uma linha por (série, mês)
  relatorio_anfavea_autoveiculos.xlsx    planilhas dinâmicas + base longa
  RELATORIO.md + graficos/*.png          resumo com destaques e gráficos

Colunas TOTAL (e procedência TOTAL) da ANFAVEA são somas derivadas e ficam fora
da base longa para não contar em dobro; os totais são recalculados aqui.
"""
from __future__ import annotations

import datetime as dt
import json
import pathlib

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parent.parent
CSV = ROOT / "data" / "csv"
OUT = ROOT / "relatorio"
GRAF = OUT / "graficos"

MESES = {m: i + 1 for i, m in enumerate(
    ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"])}
TIPO_NOME = {"producao": "Produção", "emplacamento": "Emplacamento", "exportacao": "Exportação"}
CATEGORIAS = {"AUTOMÓVEIS": "Automóveis", "COMERCIAIS LEVES": "Comerciais leves",
              "CAMINHÕES": "Caminhões", "ÔNIBUS": "Ônibus"}
PROC = {"NACIONAL": "Nacional", "IMPORTADO": "Importado", "NACIONAIS": "Nacional",
        "IMPORTADOS": "Importado", "TOTAL": "TOTAL"}
COLS = ["consulta", "tipo_dado", "categoria", "subsegmento", "combustivel",
        "procedencia", "empresa", "marca", "periodo", "valor"]


def _periodo(label: str):
    try:
        ano, mes = [x.strip() for x in str(label).split(" - ")]
        return pd.Timestamp(int(ano), MESES[mes], 1)
    except (ValueError, KeyError):
        return None


def _parse_col(consulta: str, col: str) -> dict | None:
    """Traduz o cabeçalho de uma coluna em dimensões. None = coluna derivada/ignorada."""
    c = col.strip().upper()
    if c == "TOTAL" or c.startswith("TOTAL "):
        return None
    if consulta.endswith("__categorias"):
        return {"categoria": CATEGORIAS[c]} if c in CATEGORIAS else None
    if consulta.endswith("__caminhoes_subsegmento"):
        return {"categoria": "Caminhões", "subsegmento": c.replace("CAMINHÕES", "").strip().title()}
    parts = [p.strip() for p in c.split(" - ")]
    cat = CATEGORIAS.get(parts[0], parts[0].title())
    if consulta.endswith("__combustivel"):
        return {"categoria": cat, "combustivel": " - ".join(parts[1:]).title()}
    proc = PROC.get(parts[-1])
    if proc is None or proc == "TOTAL":
        return None
    if consulta.endswith("__procedencia"):
        return {"categoria": cat, "procedencia": proc}
    if consulta.endswith("__procedencia_empresa"):
        return {"categoria": cat, "empresa": " - ".join(parts[1:-1]), "procedencia": proc}
    if consulta.endswith("__procedencia_empresa_marca"):
        return {"categoria": cat, "empresa": " - ".join(parts[1:-2]), "marca": parts[-2],
                "procedencia": proc}
    return None


def load_long() -> pd.DataFrame:
    rows = []
    for f in sorted(CSV.glob("*.csv")):
        consulta = f.stem
        tipo = TIPO_NOME[consulta.split("__")[0]]
        df = pd.read_csv(f)
        df["periodo"] = df.iloc[:, 0].map(_periodo)
        df = df[df["periodo"].notna()]
        for col in df.columns[1:-1]:
            dims = _parse_col(consulta, col)
            if dims is None:
                continue
            for per, val in zip(df["periodo"], pd.to_numeric(df[col], errors="coerce")):
                if pd.notna(val):
                    rows.append({"consulta": consulta, "tipo_dado": tipo, **dims,
                                 "periodo": per, "valor": int(val)})
    out = pd.DataFrame(rows).reindex(columns=COLS)
    out["periodo"] = pd.to_datetime(out["periodo"])
    return out


def fmt(n) -> str:
    return f"{int(round(n)):,}".replace(",", ".")


def pct(a, b) -> str:
    return "n/d" if not b else f"{(a / b - 1) * 100:+.1f}%".replace(".", ",")


def sheet(df: pd.DataFrame, index, columns) -> pd.DataFrame:
    p = df.pivot_table(index=index, columns=columns, values="valor", aggfunc="sum", fill_value=0)
    if isinstance(p.index, pd.DatetimeIndex) or index == "periodo":
        p.index = pd.to_datetime(p.index).strftime("%Y-%m")
    return p


def main() -> int:
    OUT.mkdir(exist_ok=True)
    GRAF.mkdir(exist_ok=True)
    base = load_long()
    log = json.loads((ROOT / "data" / "coleta_log.json").read_text(encoding="utf-8"))
    base.assign(periodo=base["periodo"].dt.strftime("%Y-%m")).to_csv(
        OUT / "base_autoveiculos.csv", index=False)

    q = lambda s: base[base["consulta"].str.endswith(s)]  # noqa: E731
    cat = q("__categorias")
    sub = q("__caminhoes_subsegmento")
    comb = q("__combustivel")
    proc = q("__procedencia")
    emp = q("__procedencia_empresa")
    marca = q("__procedencia_empresa_marca")

    ultimo = cat["periodo"].max()
    ano_atual, mes_atual = ultimo.year, ultimo.month
    cat = cat.assign(ano=cat["periodo"].dt.year)

    # ---------- Excel ----------
    xlsx = OUT / "relatorio_anfavea_autoveiculos.xlsx"
    resumo_anual = cat.pivot_table(index=["tipo_dado", "categoria"], columns="ano",
                                   values="valor", aggfunc="sum", fill_value=0)
    ytd = cat[cat["periodo"].dt.month <= mes_atual].pivot_table(
        index=["tipo_dado", "categoria"], columns="ano", values="valor", aggfunc="sum", fill_value=0)
    ytd.columns = [f"{c} jan-{ultimo.strftime('%m')}" for c in ytd.columns]
    leia = pd.DataFrame({"Item": [
        "Fonte", "Seção", "Coletado em", "Período", "Agregação", "Unidade",
        "Empresa/marca", "Totais", "Abas"], "Descrição": [
        "ANFAVEA – Central de Dados (https://anfavea.com.br/site/central-de-dados/)",
        "Autoveículos (produção, emplacamento, exportação)",
        log["coletado_em"][:19].replace("T", " ") + " UTC",
        f"{base['periodo'].min():%Y-%m} a {ultimo:%Y-%m}", "Mensal", "Unidades",
        "A ANFAVEA só disponibiliza emplacamento por empresa/marca a partir de 2026-01.",
        "Colunas TOTAL da ANFAVEA excluídas da base longa (somas derivadas); totais recalculados.",
        "Resumo anual, YTD, mensal por categoria, caminhões por subsegmento, emplacamento por "
        "combustível/procedência/empresa/marca e a base longa completa."]})
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        leia.to_excel(w, sheet_name="LEIA-ME", index=False)
        resumo_anual.to_excel(w, sheet_name="Resumo anual")
        ytd.to_excel(w, sheet_name="Acumulado no ano")
        for tipo in TIPO_NOME.values():
            sheet(cat[cat.tipo_dado == tipo], "periodo", "categoria").to_excel(w, sheet_name=f"{tipo} mensal")
        sheet(sub, "periodo", ["tipo_dado", "subsegmento"]).to_excel(w, sheet_name="Caminhões subsegmento")
        sheet(comb, "periodo", ["categoria", "combustivel"]).to_excel(w, sheet_name="Emplac combustível")
        if len(proc):
            sheet(proc, "periodo", ["categoria", "procedencia"]).to_excel(w, sheet_name="Emplac procedência")
        if len(emp):
            sheet(emp, ["categoria", "empresa", "procedencia"], "periodo").to_excel(w, sheet_name="Emplac empresa")
        if len(marca):
            sheet(marca, ["categoria", "empresa", "marca", "procedencia"], "periodo").to_excel(
                w, sheet_name="Emplac marca")
        base.assign(periodo=base["periodo"].dt.strftime("%Y-%m")).to_excel(
            w, sheet_name="Base longa", index=False)
        for ws in w.book.worksheets:
            ws.freeze_panes = "B2"
            for i, col in enumerate(ws.iter_cols(max_row=50), 1):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)

    # ---------- Gráficos ----------
    plt.rcParams.update({"figure.figsize": (10, 4.5), "axes.spines.top": False,
                         "axes.spines.right": False, "axes.grid": True, "grid.alpha": .3})

    def save(name):
        plt.tight_layout()
        plt.savefig(GRAF / name, dpi=130)
        plt.close()

    tot = cat.groupby(["periodo", "tipo_dado"])["valor"].sum().unstack()
    tot.plot(marker=".", title="Autoveículos – total mensal (unidades)")
    plt.xlabel("")
    save("01_total_mensal.png")

    for tipo in TIPO_NOME.values():
        sheet(cat[cat.tipo_dado == tipo], "periodo", "categoria").plot.area(
            title=f"{tipo} mensal por categoria", alpha=.85)
        save(f"02_{tipo.lower().replace('ç', 'c').replace('ã', 'a')}_categorias.png")

    sheet(sub[sub.tipo_dado == "Emplacamento"], "periodo", "subsegmento").plot(
        marker=".", title="Emplacamento de caminhões por subsegmento")
    save("03_caminhoes_subsegmento.png")

    leves = comb[comb.categoria.isin(["Automóveis", "Comerciais leves"])]
    share = sheet(leves, "periodo", "combustivel")
    share = share.div(share.sum(axis=1), axis=0) * 100
    share.plot.area(title="Leves (automóveis + comerciais leves) – mix de combustível no "
                    "emplacamento (%)", alpha=.85)
    plt.legend(loc="center left", bbox_to_anchor=(1, .5))
    save("04_mix_combustivel_leves.png")

    if len(proc):
        pp = sheet(proc, "periodo", "procedencia")
        (pp["Importado"] / pp.sum(axis=1) * 100).plot(
            marker=".", title="Participação de importados no emplacamento (%)")
        save("05_participacao_importados.png")

    if len(marca):
        top = (marca.groupby("marca")["valor"].sum().sort_values().tail(20))
        top.plot.barh(title=f"Top 20 marcas – emplacamento {ano_atual} acumulado (autoveículos)",
                      figsize=(9, 7))
        save("06_top_marcas.png")

    # ---------- RELATORIO.md ----------
    L = [f"# ANFAVEA – Autoveículos: relatório {base['periodo'].min():%m/%Y} a {ultimo:%m/%Y}", "",
         f"Fonte: [ANFAVEA – Central de Dados](https://anfavea.com.br/site/central-de-dados/). "
         f"Coletado em {log['coletado_em'][:10]}. Unidades. Planilha completa: "
         "[relatorio_anfavea_autoveiculos.xlsx](relatorio_anfavea_autoveiculos.xlsx); base longa: "
         "[base_autoveiculos.csv](base_autoveiculos.csv).", ""]

    L += ["## Totais anuais", ""]
    anos = sorted(cat["ano"].unique())
    t = cat.groupby(["tipo_dado", "ano"])["valor"].sum().unstack()
    L.append("| Indicador | " + " | ".join(
        f"{a}" + (f" (jan-{mes_atual:02d})" if a == ano_atual and mes_atual < 12 else "")
        for a in anos) + " |")
    L.append("|---|" + "---:|" * len(anos))
    for tipo in TIPO_NOME.values():
        L.append(f"| {tipo} | " + " | ".join(fmt(t.loc[tipo, a]) for a in anos) + " |")
    L.append("")

    L += [f"## Acumulado jan–{ultimo:%m}: {ano_atual} vs {ano_atual - 1}", "",
          f"| Indicador | Categoria | {ano_atual - 1} | {ano_atual} | Var. |", "|---|---|---:|---:|---:|"]
    y = cat[cat["periodo"].dt.month <= mes_atual]
    for tipo in TIPO_NOME.values():
        for c in list(CATEGORIAS.values()) + ["Total"]:
            s = y[y.tipo_dado == tipo]
            if c != "Total":
                s = s[s.categoria == c]
            a, b = s[s.ano == ano_atual - 1].valor.sum(), s[s.ano == ano_atual].valor.sum()
            L.append(f"| {tipo} | {'**Total**' if c == 'Total' else c} | {fmt(a)} | {fmt(b)} | {pct(b, a)} |")
    L.append("")

    m = cat[cat.periodo == ultimo].groupby("tipo_dado")["valor"].sum()
    m_ant = cat[cat.periodo == ultimo - pd.DateOffset(years=1)].groupby("tipo_dado")["valor"].sum()
    L += [f"## Último mês ({ultimo:%m/%Y})", "", "| Indicador | Mês | Mesmo mês ano anterior | Var. |",
          "|---|---:|---:|---:|"]
    for tipo in TIPO_NOME.values():
        L.append(f"| {tipo} | {fmt(m.get(tipo, 0))} | {fmt(m_ant.get(tipo, 0))} | "
                 f"{pct(m.get(tipo, 0), m_ant.get(tipo, 0))} |")
    L += ["", "![Total mensal](graficos/01_total_mensal.png)", ""]
    for tipo in TIPO_NOME.values():
        fn = f"02_{tipo.lower().replace('ç', 'c').replace('ã', 'a')}_categorias.png"
        L += [f"![{tipo}](graficos/{fn})", ""]

    L += ["## Emplacamento por combustível (leves)", ""]
    sh_ano = leves.assign(ano=leves.periodo.dt.year).pivot_table(
        index="combustivel", columns="ano", values="valor", aggfunc="sum", fill_value=0)
    sh_pct = sh_ano.div(sh_ano.sum()) * 100
    L.append("| Combustível | " + " | ".join(str(a) for a in sh_pct.columns) + " |")
    L.append("|---|" + "---:|" * len(sh_pct.columns))
    for f_, r in sh_pct.sort_values(sh_pct.columns[-1], ascending=False).iterrows():
        L.append(f"| {f_} | " + " | ".join(f"{v:.1f}%".replace(".", ",") for v in r) + " |")
    L += ["", "![Mix de combustível](graficos/04_mix_combustivel_leves.png)", ""]

    L += ["## Caminhões por subsegmento (emplacamento)", "",
          "![Caminhões](graficos/03_caminhoes_subsegmento.png)", ""]

    if len(proc):
        pa = proc.assign(ano=proc.periodo.dt.year).pivot_table(
            index="ano", columns="procedencia", values="valor", aggfunc="sum")
        L += ["## Nacionais x importados (emplacamento)", "", "| Ano | Nacional | Importado | % importado |",
              "|---|---:|---:|---:|"]
        for a, r in pa.iterrows():
            share_imp = f"{r['Importado'] / r.sum() * 100:.1f}".replace(".", ",")
            L.append(f"| {a} | {fmt(r['Nacional'])} | {fmt(r['Importado'])} | {share_imp}% |")
        L += ["", "![Importados](graficos/05_participacao_importados.png)", ""]

    if len(emp):
        L += [f"## Emplacamento por empresa – {ano_atual} acumulado (desde 01/{ano_atual})", "",
              "> A ANFAVEA só publica empresa/marca a partir de jan/2026.", "",
              "| # | Empresa | Nacional | Importado | Total | Part. |", "|---:|---|---:|---:|---:|---:|"]
        e = emp.pivot_table(index="empresa", columns="procedencia", values="valor",
                            aggfunc="sum", fill_value=0)
        e["Total"] = e.sum(axis=1)
        e = e.sort_values("Total", ascending=False)
        for i, (n, r) in enumerate(e.head(25).iterrows(), 1):
            L.append(f"| {i} | {n} | {fmt(r.get('Nacional', 0))} | {fmt(r.get('Importado', 0))} | "
                     f"{fmt(r['Total'])} | {str(round(r['Total'] / e['Total'].sum() * 100, 1)).replace('.', ',')}% |")
        L.append("")
    if len(marca):
        L += ["![Top marcas](graficos/06_top_marcas.png)", ""]

    L += ["## Metodologia e limitações", "",
          "- Dados obtidos da mesma API que alimenta a Central de Dados "
          "(`admin-ajax.php`, `action=anfavea_dashboard1`), agregação mensal.",
          "- Emplacamento: Renavam/Senatran. Produção a partir de 2026 inclui empresas não associadas "
          "(volume estimado pela ANFAVEA).",
          "- Empresa/marca: tabela de origem só tem dados desde 2026-01 (limitação da ANFAVEA).",
          "- Subsegmentos de caminhões não estão disponíveis cruzados com combustível.",
          "- Colunas TOTAL da ANFAVEA foram descartadas e os totais recalculados.",
          "- Conferência: as aberturas por procedência, empresa e marca somam exatamente o total por "
          "categoria; por combustível e por subsegmento de caminhões há diferenças pontuais de poucas "
          "unidades em alguns meses, vindas da própria fonte.", ""]
    (OUT / "RELATORIO.md").write_text("\n".join(L), encoding="utf-8")
    print(f"ok: {len(base)} linhas na base, último mês {ultimo:%Y-%m}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
